"""调研报告生成器 v1：只生成第二章「调研过程」。

链路强制：调研事件 → LLM 正文带 ``[EV:{id}]`` 引用标记 → 图注/图片 → 材料溯源表。
LLM 未配置或调用失败时使用确定性模板（同样带引用标记），保证离线可用。
"""
from __future__ import annotations

import asyncio
import json
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from sqlalchemy.orm import Session

from ..core.config import settings
from ..core.database import SessionLocal
from ..models.survey import SurveyEvent, SurveyReport, SurveyTask
from ..models.track import TrackFile, TrackPhoto
from .llm_provider import get_llm_provider

CHAPTER_KEY = "chapter2"
CHAPTER_TITLE = "二、调研过程"
EV_RE = re.compile(r"\[EV:(\d+)\]")
MAX_FIGURES = 8


def _event_line(event: SurveyEvent) -> str:
    time_text = event.timestamp.strftime("%Y-%m-%d %H:%M") if event.timestamp else "时间未知"
    address = event.address or "位置未知"
    return f"[EV:{event.id}] {time_text} {address} 《{event.title or '调研记录'}》：{event.content or ''}"


def _build_prompt(task: SurveyTask, events: list[SurveyEvent], stats: dict) -> str:
    event_lines = "\n".join(_event_line(event) for event in events) or "（无可用事件）"
    return f"""你是文化遗产调研报告撰写助手。以下是本次田野调研的事件清单（按时间排序）与轨迹概况。

请撰写报告「二、调研过程」章节正文，包含三部分：
1. time：调研时间（依据轨迹起止时间客观陈述）；
2. method：调研方式（写明：实地踏勘，同步记录 GPS 轨迹并采集现场影像）；
3. route：实地踏勘过程，按时间顺序叙述踏勘路线与关键点位。

硬性要求：
- 学术、客观口径；不得编造事件清单之外的地点、数字或结论；
- 每一个事实性句子都必须带引用标记 [EV:事件ID]，只能引用下列事件 ID；
- 输出严格 JSON（不要 markdown 代码块）：{{"time": "...", "method": "...", "route": "..."}}

任务标题：{task.title}
轨迹概况：轨迹点 {stats.get('track_point_count', 0)} 个，照片 {stats.get('photo_count', 0)} 张，总里程约 {stats.get('distance_km', 0)} km，起止时间 {stats.get('start_time') or '未知'} ~ {stats.get('end_time') or '未知'}。

事件清单：
{event_lines}"""


def _parse_llm_json(text: str) -> dict | None:
    cleaned = (text or "").strip()
    if not cleaned:
        return None
    cleaned = re.sub(r"^```[a-zA-Z]*\s*", "", cleaned)
    cleaned = re.sub(r"\s*```$", "", cleaned).strip()
    for candidate in (cleaned, (re.search(r"\{[\s\S]*\}", cleaned).group(0) if re.search(r"\{[\s\S]*\}", cleaned) else "")):
        if not candidate:
            continue
        try:
            data = json.loads(candidate)
            if isinstance(data, dict) and data.get("route"):
                return data
        except Exception:
            continue
    return None


async def _call_llm(prompt: str) -> dict | None:
    try:
        provider = get_llm_provider()
        raw = await provider.chat([
            {"role": "system", "content": "你是文化遗产调研报告撰写助手，只输出要求格式的内容。"},
            {"role": "user", "content": prompt},
        ])
    except Exception:
        return None
    return _parse_llm_json(raw)


def _fallback_texts(events: list[SurveyEvent]) -> dict:
    if not events:
        return {
            "time": "本次调研暂无可用于报告的轨迹或影像事件。",
            "method": "调研方式为实地踏勘，轨迹记录与影像采集。",
            "route": "本次调研暂无可用于报告的调研事件。",
        }
    order = sorted(events, key=lambda item: (item.timestamp is None, item.timestamp or 0))
    first, last = order[0], order[-1]
    start = first.timestamp.strftime("%Y-%m-%d %H:%M") if first.timestamp else "时间未知"
    end = last.timestamp.strftime("%Y-%m-%d %H:%M") if last.timestamp else "时间未知"
    sentences = []
    for event in order:
        time_text = event.timestamp.strftime("%H:%M") if event.timestamp else "时间未知"
        place = event.address or "调研点位"
        title = event.title or "调研记录"
        sentences.append(f"{time_text}，调研人员到达{place}，记录「{title}」[EV:{event.id}]。")
    return {
        "time": f"本次调研时间为 {start} 至 {end}。",
        "method": "本次调研采用实地踏勘方式，同步记录 GPS 轨迹并采集现场影像资料。",
        "route": "".join(sentences),
    }


def _split_paragraphs(text: str) -> list[str]:
    parts = [part.strip() for part in re.split(r"\n\s*\n", (text or "").strip()) if part.strip()]
    if len(parts) <= 1:
        single = (text or "").strip()
        if not single:
            return []
        parts = [part.strip() for part in single.split("\n") if part.strip()]
    return parts


def _figure_of(photo: TrackPhoto) -> dict:
    path = photo.file_path if photo.file_path and Path(photo.file_path).exists() else photo.thumb_path
    return {
        "photo_id": photo.id,
        "path": path,
        "caption": photo.caption or photo.file_name or "调研影像",
        "shot_time": photo.shot_time.strftime("%Y-%m-%d %H:%M") if photo.shot_time else "",
        "address": photo.address or "",
    }


def _pick_representatives(cited_ids: list[int], candidates: list[tuple[SurveyEvent, TrackPhoto]]) -> list[tuple[SurveyEvent, TrackPhoto]]:
    selected: list[tuple[SurveyEvent, TrackPhoto]] = []
    chosen: set[int] = set()
    by_event = {event.id: (event, photo) for event, photo in candidates}

    for event_id in cited_ids:
        if len(selected) >= MAX_FIGURES:
            break
        if event_id in by_event and event_id not in chosen:
            selected.append(by_event[event_id])
            chosen.add(event_id)

    remaining = [item for item in candidates if item[0].id not in chosen]
    slots = MAX_FIGURES - len(selected)
    if slots > 0 and remaining:
        if len(remaining) <= slots:
            selected.extend(remaining)
        else:
            step = len(remaining) / slots
            for index in range(slots):
                selected.append(remaining[min(len(remaining) - 1, int(index * step))])
    return selected


def _build_citations(chapter: dict, figures: list[dict]) -> list[dict]:
    citations: list[dict] = []
    seen: set[int] = set()
    for section in chapter["sections"]:
        for block in section["blocks"]:
            if block["type"] == "text":
                for event_id in EV_RE.findall(block["text"]):
                    eid = int(event_id)
                    if eid in seen:
                        continue
                    seen.add(eid)
                    citations.append({"event_id": eid, "section": section["heading"], "role": "正文引用"})
    for figure in figures:
        eid = figure.get("event_id")
        if eid is None or eid in seen:
            continue
        seen.add(eid)
        citations.append({"event_id": eid, "section": figure["section"], "role": "配图"})
    return citations


def _decorate_citations(db: Session, citations: list[dict]) -> list[dict]:
    ids = [item["event_id"] for item in citations]
    events = db.query(SurveyEvent).filter(SurveyEvent.id.in_(ids)).all() if ids else []
    mapping = {event.id: event for event in events}
    for item in citations:
        event = mapping.get(item["event_id"])
        item["title"] = (event.title or "") if event else ""
        item["time"] = event.timestamp.strftime("%Y-%m-%d %H:%M") if event and event.timestamp else ""
        item["address"] = (event.address or "") if event else ""
    return citations


async def _build_chapter(db: Session, task: SurveyTask) -> dict:
    events = (
        db.query(SurveyEvent)
        .filter(SurveyEvent.task_id == task.id, SurveyEvent.usable_for_report.is_(True))
        .order_by(SurveyEvent.timestamp.asc().nullslast(), SurveyEvent.id.asc())
        .all()
    )
    track = (
        db.query(TrackFile)
        .filter(TrackFile.survey_task_id == task.id)
        .order_by(TrackFile.id.desc())
        .first()
    )
    photos = {}
    if track is not None:
        photos = {photo.id: photo for photo in db.query(TrackPhoto).filter(TrackPhoto.track_id == track.id).all()}

    stats = {
        "track_point_count": track.track_point_count if track else 0,
        "photo_count": track.photo_count if track else 0,
        "distance_km": round((track.distance_meters or 0) / 1000.0, 2) if track else 0,
        "start_time": track.start_time.strftime("%Y-%m-%d %H:%M") if track and track.start_time else None,
        "end_time": track.end_time.strftime("%Y-%m-%d %H:%M") if track and track.end_time else None,
    }

    texts = await _call_llm(_build_prompt(task, events, stats)) if events else None
    if texts and events:
        # 只允许引用真实存在的事件 ID；无有效引用标记则整段回落到确定性模板
        valid_ids = {event.id for event in events}
        texts = {
            key: EV_RE.sub(lambda match: match.group(0) if int(match.group(1)) in valid_ids else "", value or "")
            for key, value in texts.items()
        }
        if not EV_RE.search(texts.get("route", "")):
            texts = _fallback_texts(events)
    if not texts:
        texts = _fallback_texts(events)

    # 组装文本块
    text_sections: list[dict] = []
    all_cited: list[int] = []
    for key, heading, body in (
        ("time", "（一）调研时间", texts.get("time", "")),
        ("method", "（二）调研方式", texts.get("method", "")),
        ("route", "（三）实地踏勘", texts.get("route", "")),
    ):
        blocks = [{"type": "text", "text": paragraph} for paragraph in _split_paragraphs(body)]
        for block in blocks:
            all_cited.extend(int(eid) for eid in EV_RE.findall(block["text"]))
        text_sections.append({"key": key, "heading": heading, "blocks": blocks})

    # 挑选代表性照片
    context = events
    candidates = [
        (event, photos.get(event.source_material_id))
        for event in context
        if event.source_material_id and photos.get(event.source_material_id) and _figure_of(photos[event.source_material_id])["path"]
    ]
    selected = _pick_representatives(all_cited, candidates)
    pending = {event.id: _figure_of(photo) | {"event_id": event.id} for event, photo in selected}

    # 图片就近插入到引用了该事件的段落之后
    route_section = next(section for section in text_sections if section["key"] == "route")
    inserted: set[int] = set()
    route_blocks: list[dict] = []
    for block in route_section["blocks"]:
        route_blocks.append(block)
        for eid in EV_RE.findall(block["text"]):
            event_id = int(eid)
            if event_id in pending and event_id not in inserted:
                figure = pending[event_id]
                route_blocks.append({"type": "figure", "photo_id": figure["photo_id"], "event_id": event_id})
                inserted.add(event_id)
    route_section["blocks"] = route_blocks

    leftover = [figure for figure in pending.values() if figure["event_id"] not in inserted]
    if leftover:
        text_sections.append(
            {
                "key": "photos",
                "heading": "（四）踏勘照片",
                "blocks": [
                    {"type": "figure", "photo_id": figure["photo_id"], "event_id": figure["event_id"]}
                    for figure in leftover
                ],
            }
        )

    # 统一编号
    figures: list[dict] = []
    counter = 0
    figure_section: dict[int, str] = {}
    for section in text_sections:
        for block in section["blocks"]:
            if block["type"] != "figure":
                continue
            counter += 1
            figure = dict(pending[block["event_id"]])
            figure["no"] = f"图 2.{counter}"
            figure["section"] = section["heading"]
            block["no"] = figure["no"]
            figures.append(figure)
            figure_section[figure["photo_id"]] = section["heading"]

    chapter = {
        "chapter": CHAPTER_KEY,
        "title": CHAPTER_TITLE,
        "stats": stats,
        "sections": text_sections,
        "figures": figures,
    }
    citations = _build_citations(chapter, figures)
    chapter["citations"] = _decorate_citations(db, citations)
    return chapter


def _set_run_font(run, *, name: str = "仿宋", size: int = 12, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    try:
        run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    except Exception:
        pass


def _add_heading(doc: Document, text: str, *, size: int = 14) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(6)
    _set_run_font(paragraph.add_run(text), name="黑体", size=size, bold=True)


def render_chapter_docx(chapter: dict, out_path: Path) -> Path:
    """按 docs/08 第 5.2 节默认模板渲染第二章 docx。"""
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "仿宋"
    normal.font.size = Pt(12)
    try:
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "仿宋")
    except Exception:
        pass

    _add_heading(document, chapter["title"], size=16)

    figures_by_photo = {figure["photo_id"]: figure for figure in chapter.get("figures", [])}

    for section in chapter["sections"]:
        _add_heading(document, section["heading"], size=14)
        for block in section["blocks"]:
            if block["type"] == "text":
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.first_line_indent = Pt(24)
                paragraph.paragraph_format.line_spacing = 1.5
                _set_run_font(paragraph.add_run(block["text"]))
            elif block["type"] == "figure":
                figure = figures_by_photo.get(block.get("photo_id"))
                path = figure.get("path") if figure else None
                if not figure or not path or not Path(path).exists():
                    continue
                picture = document.add_paragraph()
                picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
                try:
                    picture.add_run().add_picture(str(path), width=Cm(14))
                except Exception:
                    continue
                caption = document.add_paragraph()
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption_text = (
                    f"{figure['no']}：{figure['caption']}"
                    f"（拍摄时间：{figure['shot_time'] or '未知'}，{figure['address'] or '位置未知'}）"
                    "（来源：本次调研拍摄）"
                )
                _set_run_font(caption.add_run(caption_text), size=10.5)

    _add_heading(document, "材料溯源表", size=16)
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["章节", "引用事件", "时间", "地点"]
    for index, header in enumerate(headers):
        _set_run_font(table.rows[0].cells[index].paragraphs[0].add_run(header), size=10.5, bold=True)
    for item in chapter.get("citations", []):
        row = table.add_row().cells
        values = [
            f"{item.get('section', '')}（{item.get('role', '')}）",
            f"EV:{item.get('event_id')} {item.get('title', '')}",
            item.get("time", ""),
            item.get("address", ""),
        ]
        for index, value in enumerate(values):
            _set_run_font(row[index].paragraphs[0].add_run(str(value)), size=10.5)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out_path))
    return out_path


def report_path(task_id: int) -> Path:
    return settings.survey_report_dir / str(task_id) / f"{CHAPTER_KEY}.docx"


def delete_task_reports(task_id: int) -> None:
    """删除某调研任务已导出的报告文件。"""
    path = settings.survey_report_dir / str(task_id)
    if path.exists():
        shutil.rmtree(path, ignore_errors=True)


def recover_interrupted_reports() -> int:
    """服务启动时把因重启而中断生成的报告标记为 failed，供用户重新生成。"""
    db = SessionLocal()
    try:
        reports = db.query(SurveyReport).filter(SurveyReport.status == "generating").all()
        for report in reports:
            report.status = "failed"
            report.error_message = "服务重启导致生成中断，请重新生成"
        db.commit()
        return len(reports)
    except Exception:
        db.rollback()
        return 0
    finally:
        db.close()


def get_or_create_report(db: Session, task_id: int) -> SurveyReport:
    report = (
        db.query(SurveyReport)
        .filter(SurveyReport.task_id == task_id, SurveyReport.chapter == CHAPTER_KEY)
        .first()
    )
    if report is None:
        report = SurveyReport(task_id=task_id, chapter=CHAPTER_KEY, version=0, status="draft")
        db.add(report)
        db.commit()
        db.refresh(report)
    return report


async def run_chapter2_generation(task_id: int) -> None:
    """后台生成/重新生成章节，结果与引用清单落库。"""
    db = SessionLocal()
    try:
        task = db.get(SurveyTask, task_id)
        report = get_or_create_report(db, task_id)
        if task is None:
            report.status = "failed"
            report.error_message = "调研任务不存在"
            db.commit()
            return
        report.status = "generating"
        report.error_message = None
        db.commit()
        try:
            chapter = await _build_chapter(db, task)
            out_path = await asyncio.to_thread(render_chapter_docx, chapter, report_path(task_id))
            report.content = json.dumps(chapter, ensure_ascii=False)
            report.citations = json.dumps(chapter.get("citations", []), ensure_ascii=False)
            report.docx_path = str(out_path)
            report.version = (report.version or 0) + 1
            report.status = "done"
        except Exception as exc:  # noqa: BLE001
            report.status = "failed"
            report.error_message = str(exc)[:2000]
        task.status = "ready" if report.status == "done" else task.status
        db.commit()
    finally:
        db.close()


def report_payload(report: SurveyReport | None) -> dict:
    if report is None:
        return {"status": "empty", "version": 0, "content": None, "citations": [], "error_message": None}
    try:
        content = json.loads(report.content) if report.content else None
    except Exception:
        content = None
    try:
        citations = json.loads(report.citations) if report.citations else []
    except Exception:
        citations = []
    return {
        "status": report.status,
        "version": report.version,
        "content": content,
        "citations": citations,
        "docx_ready": bool(report.docx_path and Path(report.docx_path).exists()),
        "error_message": report.error_message,
        "updated_at": report.updated_at.isoformat() if report.updated_at else None,
    }
