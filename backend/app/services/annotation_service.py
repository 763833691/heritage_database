"""照片语义标注工作台业务逻辑：样本导入、盲标取题、编码提交、信度计算与导出。"""
from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path
from typing import Optional

from sqlalchemy.orm import Session

from ..models.annotation import AnnotationItem, AnnotationRecord, AnnotationTask

_BACKEND_DIR = Path(__file__).resolve().parents[2]
_DEFAULT_SAMPLE = "data/annotation/sample_v1.json"

# A–G 编码类目（与编码手册一致）
CATEGORY_NAMES: dict[str, str] = {
    "A": "遗址本体展示",
    "B": "复原建筑与覆罩",
    "C": "阐释解说设施",
    "D": "数字互动展示",
    "E": "运营与消费场景",
    "F": "景观环境与城市关系",
    "G": "其他",
}
CATEGORY_KEYS = tuple(CATEGORY_NAMES.keys())

PLATFORM_KEY = "platform"
AI_KEY = "ai"


def photo_url(track_id: Optional[int], photo_id: Optional[int]) -> str:
    return f"/api/track/{track_id}/photos/{photo_id}"


def cohens_kappa(codes_a: list[str], codes_b: list[str]) -> float:
    """Cohen's Kappa：仅统计双方都非空的配对；无有效配对或 pe==1 时返回 0.0。"""
    pairs = [(a, b) for a, b in zip(codes_a, codes_b) if a and b]
    n = len(pairs)
    if n == 0:
        return 0.0
    po = sum(1 for a, b in pairs if a == b) / n
    labels = set()
    for a, b in pairs:
        labels.add(a)
        labels.add(b)
    pe = 0.0
    for label in labels:
        count_a = sum(1 for a, _ in pairs if a == label) / n
        count_b = sum(1 for _, b in pairs if b == label) / n
        pe += count_a * count_b
    if pe >= 1.0:
        return 0.0
    return (po - pe) / (1 - pe)


def _resolve_sample_path(sample_path: Optional[str]) -> Path:
    raw = sample_path or _DEFAULT_SAMPLE
    path = Path(raw)
    if not path.is_absolute():
        path = _BACKEND_DIR / raw
    return path


def _task_items(db: Session, task_id: int) -> list[AnnotationItem]:
    return (
        db.query(AnnotationItem)
        .filter(AnnotationItem.task_id == task_id)
        .order_by(AnnotationItem.seq.asc(), AnnotationItem.id.asc())
        .all()
    )


def _task_coders(db: Session, task_id: int) -> list[str]:
    rows = (
        db.query(AnnotationRecord.coder)
        .join(AnnotationItem, AnnotationRecord.item_id == AnnotationItem.id)
        .filter(AnnotationItem.task_id == task_id)
        .distinct()
        .all()
    )
    return sorted({row[0] for row in rows if row[0]})


def _task_summary(db: Session, task: AnnotationTask) -> dict:
    total = db.query(AnnotationItem).filter(AnnotationItem.task_id == task.id).count()
    counts = (
        db.query(AnnotationRecord.coder, AnnotationRecord.item_id)
        .join(AnnotationItem, AnnotationRecord.item_id == AnnotationItem.id)
        .filter(AnnotationItem.task_id == task.id)
        .all()
    )
    done: dict[str, int] = {}
    for coder, _item_id in counts:
        done[coder] = done.get(coder, 0) + 1
    coders = [{"coder": coder, "done": count} for coder, count in sorted(done.items())]
    return {
        "id": task.id,
        "name": task.name,
        "codebook_version": task.codebook_version,
        "status": task.status,
        "created_at": task.created_at.isoformat() if task.created_at else None,
        "total": total,
        "coders": coders,
        "categories": CATEGORY_NAMES,
    }


def import_task(db: Session, name: Optional[str], sample_path: Optional[str]) -> tuple[AnnotationTask, bool]:
    """从样本 JSON 导入任务；同名任务已存在时直接返回既有任务（幂等）。"""
    path = _resolve_sample_path(sample_path)
    if not path.exists():
        raise FileNotFoundError(f"样本文件不存在：{path}")
    data = json.loads(path.read_text(encoding="utf-8"))

    task_name = (name or data.get("task_name") or "标注任务").strip()
    existing = db.query(AnnotationTask).filter(AnnotationTask.name == task_name).first()
    if existing is not None:
        return existing, False

    task = AnnotationTask(
        name=task_name,
        codebook_version=data.get("codebook_version") or "v1.0",
    )
    db.add(task)
    db.flush()

    for index, entry in enumerate(data.get("items", [])):
        db.add(
            AnnotationItem(
                task_id=task.id,
                sid=entry.get("sid"),
                track_id=entry.get("track_id"),
                photo_id=entry.get("photo_id"),
                seq=index,  # 任务内展示顺序：按样本文件中的条目顺序
                platform_type=(entry.get("platform_type") or None),
                ai_code=(entry.get("ai_code") or None),
                caption=entry.get("caption"),
            )
        )
    db.commit()
    db.refresh(task)
    return task, True


def list_tasks(db: Session) -> list[dict]:
    tasks = db.query(AnnotationTask).order_by(AnnotationTask.id.desc()).all()
    return [_task_summary(db, task) for task in tasks]


def get_task(db: Session, task_id: int) -> Optional[AnnotationTask]:
    return db.get(AnnotationTask, task_id)


def _progress(db: Session, task_id: int, coder: str) -> dict:
    total = db.query(AnnotationItem).filter(AnnotationItem.task_id == task_id).count()
    done = (
        db.query(AnnotationRecord)
        .join(AnnotationItem, AnnotationRecord.item_id == AnnotationItem.id)
        .filter(AnnotationItem.task_id == task_id, AnnotationRecord.coder == coder)
        .count()
    )
    return {"done": done, "total": total}


def next_item(db: Session, task_id: int, coder: str) -> dict:
    """取该编码者尚未标注的第一张（按 seq 排序）；全部完成返回 finished。"""
    progress = _progress(db, task_id, coder)
    coded = db.query(AnnotationRecord.item_id).filter(AnnotationRecord.coder == coder)
    item = (
        db.query(AnnotationItem)
        .filter(AnnotationItem.task_id == task_id, ~AnnotationItem.id.in_(coded))
        .order_by(AnnotationItem.seq.asc(), AnnotationItem.id.asc())
        .first()
    )
    if item is None:
        return {"finished": True, "progress": progress}
    return {
        "item_id": item.id,
        "sid": item.sid,
        "photo_url": photo_url(item.track_id, item.photo_id),
        "progress": progress,
    }


def photo_meta(db: Session, item_id: int) -> dict:
    item = db.get(AnnotationItem, item_id)
    if item is None:
        raise ValueError("条目不存在")
    return {"sid": item.sid, "photo_url": photo_url(item.track_id, item.photo_id)}


def submit_code(db: Session, item_id: int, coder: str, code: str, note: str = "") -> AnnotationRecord:
    code = (code or "").strip().upper()
    if code not in CATEGORY_NAMES:
        raise ValueError("code 应为 A–G 之一")
    coder = (coder or "").strip()
    if not coder:
        raise ValueError("coder 不能为空")
    item = db.get(AnnotationItem, item_id)
    if item is None:
        raise ValueError("条目不存在")

    record = (
        db.query(AnnotationRecord)
        .filter(AnnotationRecord.item_id == item_id, AnnotationRecord.coder == coder)
        .first()
    )
    if record is None:
        record = AnnotationRecord(item_id=item_id, coder=coder, code=code, note=note or "")
        db.add(record)
    else:
        record.code = code
        record.note = note or ""
    db.commit()
    db.refresh(record)
    return record


def _pairwise(items: list[AnnotationItem], coder_codes: dict[int, dict[str, str]], a_key: str, b_key: str) -> dict:
    def value(item: AnnotationItem, key: str) -> Optional[str]:
        if key == PLATFORM_KEY:
            return item.platform_type or None
        if key == AI_KEY:
            return item.ai_code or None
        return coder_codes.get(item.id, {}).get(key) or None

    codes_a: list[str] = []
    codes_b: list[str] = []
    for item in items:
        av = value(item, a_key)
        bv = value(item, b_key)
        if av and bv:
            codes_a.append(av)
            codes_b.append(bv)

    n = len(codes_a)
    agree = sum(1 for a, b in zip(codes_a, codes_b) if a == b)
    return {
        "a": a_key,
        "b": b_key,
        "n": n,
        "agree": agree,
        "agree_rate": round(agree / n, 3) if n else 0.0,
        "kappa": round(cohens_kappa(codes_a, codes_b), 3),
    }


def results(db: Session, task_id: int) -> dict:
    task = db.get(AnnotationTask, task_id)
    if task is None:
        raise ValueError("任务不存在")
    items = _task_items(db, task_id)
    coders = _task_coders(db, task_id)

    records = (
        db.query(AnnotationRecord)
        .join(AnnotationItem, AnnotationRecord.item_id == AnnotationItem.id)
        .filter(AnnotationItem.task_id == task_id)
        .all()
    )
    coder_codes: dict[int, dict[str, str]] = {}
    coder_notes: dict[int, dict[str, str]] = {}
    for record in records:
        coder_codes.setdefault(record.item_id, {})[record.coder] = record.code
        if record.note:
            coder_notes.setdefault(record.item_id, {})[record.coder] = record.note

    pairwise: list[dict] = [_pairwise(items, coder_codes, PLATFORM_KEY, AI_KEY)]
    for coder in coders:
        pairwise.append(_pairwise(items, coder_codes, coder, AI_KEY))
        pairwise.append(_pairwise(items, coder_codes, coder, PLATFORM_KEY))
    for i in range(len(coders)):
        for j in range(i + 1, len(coders)):
            pairwise.append(_pairwise(items, coder_codes, coders[i], coders[j]))

    disagreements: list[dict] = []
    for item in items:
        human = {coder: coder_codes.get(item.id, {}).get(coder) for coder in coders}
        human = {k: v for k, v in human.items() if v}
        values = [v for v in [item.platform_type, item.ai_code, *human.values()] if v]
        if len(set(values)) > 1:
            disagreements.append(
                {
                    "item_id": item.id,
                    "sid": item.sid,
                    "photo_url": photo_url(item.track_id, item.photo_id),
                    "platform": item.platform_type,
                    "ai": item.ai_code,
                    "human": human,
                    "notes": coder_notes.get(item.id, {}),
                    "caption": item.caption,
                }
            )

    return {
        "task": {"id": task.id, "name": task.name, "total": len(items)},
        "coders": coders,
        "pairwise": pairwise,
        "disagreements": disagreements,
    }


def build_xlsx(db: Session, task_id: int) -> bytes:
    """导出标注结果 Excel：每行一个样本，含平台/AI/各人工编码、是否一致与备注。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    task = db.get(AnnotationTask, task_id)
    if task is None:
        raise ValueError("任务不存在")
    items = _task_items(db, task_id)
    coders = _task_coders(db, task_id)

    records = (
        db.query(AnnotationRecord)
        .join(AnnotationItem, AnnotationRecord.item_id == AnnotationItem.id)
        .filter(AnnotationItem.task_id == task_id)
        .all()
    )
    coder_codes: dict[int, dict[str, str]] = {}
    coder_notes: dict[int, dict[str, str]] = {}
    for record in records:
        coder_codes.setdefault(record.item_id, {})[record.coder] = record.code
        if record.note:
            coder_notes.setdefault(record.item_id, {})[record.coder] = record.note

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "标注结果"

    headers = ["样本编号", "轨迹ID", "照片ID", "平台分类", "智能体编码"]
    headers += [f"人工-{coder}" for coder in coders]
    headers += ["是否一致", "备注"]
    sheet.append(headers)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="085399")
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill

    for item in items:
        human = {coder: coder_codes.get(item.id, {}).get(coder) for coder in coders}
        values = [v for v in [item.platform_type, item.ai_code, *human.values()] if v]
        consistent = "是" if len(set(values)) <= 1 else "否"
        notes = coder_notes.get(item.id, {})
        note_text = "; ".join(f"{coder}:{note}" for coder, note in notes.items())
        row = [item.sid, item.track_id, item.photo_id, item.platform_type, item.ai_code]
        row += [human.get(coder) or "" for coder in coders]
        row += [consistent, note_text]
        sheet.append(row)

    widths = [12, 10, 10, 10, 12] + [12] * len(coders) + [10, 40]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
