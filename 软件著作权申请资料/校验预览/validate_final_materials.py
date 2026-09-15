from __future__ import annotations

import json
from pathlib import Path

from docx import Document


ROOT = Path(r"H:\my_code\datak")
WORKDIR = ROOT / "软件著作权申请资料"
FINAL_DIR = WORKDIR / "正式资料"
SOFTWARE_NAME = "考古遗址公园文化价值智能数字化阐释平台"
VERSION = "V1.0"
COPYRIGHT_HOLDERS = ["王新文", "刘泽宇", "张棪森", "梅子阳", "邱勇杰"]


def header_text(document: Document) -> str:
    chunks: list[str] = []
    for section in document.sections:
        chunks.extend(p.text for p in section.header.paragraphs)
        for table in section.header.tables:
            chunks.extend(
                cell.text for row in table.rows for cell in row.cells
            )
    return " ".join(chunks)


def main() -> None:
    selection = json.loads(
        (WORKDIR / "草稿" / "代码文件选择.json").read_text(encoding="utf-8")
    )
    selected = [
        item["path"] for item in selection["files"] if item.get("selected")
    ]
    source_lines: set[str] = set()
    for relative_path in selected:
        source_text = (ROOT / relative_path).read_text(
            encoding="utf-8", errors="replace"
        )
        source_lines.update(line.rstrip("\r") for line in source_text.splitlines())

    code_results = []
    for filename in [
        f"{SOFTWARE_NAME}-代码(前30页).docx",
        f"{SOFTWARE_NAME}-代码(后30页).docx",
    ]:
        document = Document(FINAL_DIR / filename)
        material_lines = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
            and not paragraph.text.startswith("// File:")
        ]
        unmatched = [line for line in material_lines if line not in source_lines]
        code_results.append(
            {
                "file": filename,
                "material_lines": len(material_lines),
                "exact_source_matches": len(material_lines) - len(unmatched),
                "unmatched": len(unmatched),
                "header_name_version_ok": (
                    f"{SOFTWARE_NAME} {VERSION}" in header_text(document)
                ),
            }
        )

    manual = Document(FINAL_DIR / f"{SOFTWARE_NAME}_操作手册.docx")
    manual_text = "\n".join(p.text for p in manual.paragraphs)
    application_text = (FINAL_DIR / "申请表信息.txt").read_text(
        encoding="utf-8"
    )
    page_report = json.loads(
        (WORKDIR / "校验预览" / "代码物理分页校验.json").read_text(
            encoding="utf-8"
        )
    )

    summary = {
        "selected_source_files": len(selected),
        "code_documents": code_results,
        "code_page_counts_ok": (
            page_report["front_rendered_pages"] == 30
            and page_report["back_rendered_pages"] == 30
        ),
        "manual_inline_images": len(manual.inline_shapes),
        "manual_name_ok": SOFTWARE_NAME in manual_text,
        "manual_header_ok": (
            f"{SOFTWARE_NAME} {VERSION}" in header_text(manual)
        ),
        "application_name_ok": SOFTWARE_NAME in application_text,
        "application_version_ok": VERSION in application_text,
        "application_names_ok": all(
            name in application_text for name in COPYRIGHT_HOLDERS
        ),
        "pending_markers": sum(
            application_text.count(marker)
            for marker in ["待确认", "待用户确认", "TODO", "TBD"]
        ),
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
